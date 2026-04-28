#!/usr/bin/env python3
"""
Build an Excel workbook for pairwise human evaluation of reasoning traces.

Default behavior:
  - A is always keentomato/harpo_hier_step400
  - B is sampled from each baseline model separately
  - samples are matched by (dataset, _orig_idx)
  - one worksheet is created per baseline model

Example:
    python build_reasoning_annotation_workbook.py

    python build_reasoning_annotation_workbook.py \
        --results_dir /scratch/keane/results/thinking_w_params \
        --samples_per_baseline 50 \
        --seed 42
"""

from __future__ import annotations

import argparse
import json
import random
from collections import defaultdict
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


RESULTS_DIR = Path("/home/keaneong/human-behavior/verl/zero_shot_inference/results/thinking_w_params")
ANCHOR_MODEL = "keentomato/harpo_hier_step400"
BASELINE_MODELS = [
    # "PhilipC/HumanOmniV2",
    # "ddvd233/OmniSapiens-7B-RL",
    # "Qwen/Qwen2.5-Omni-7B",
    "google_gemma-4-e4b-it"
]
KNOWN_DATASETS = [
    "av-asd",
    "iemocap",
    "eatd",
    "mvsa",
    # "dreaddit",
    "sarcnet",
]
DEFAULT_SAMPLES_PER_BASELINE = 50
DEFAULT_SEED = 42
EXCEL_MAX_CELL_LENGTH = 32767

# W&B config. Set WANDB_PROJECT="" to disable artifact upload.
WANDB_PROJECT = "reasoning-evaluation"
WANDB_RUN_ID = "reasoning_annotation_workbook"
WANDB_RUN_NAME = "reasoning_annotation_workbook"
WANDB_ENTITY = ""
WANDB_ARTIFACT_NAME = "reasoning_annotation_workbook"
WANDB_ARTIFACT_TYPE = "annotation_workbook"


def slugify_model(model_name: str) -> str:
    return model_name.replace("/", "_")


def detect_dataset_from_filename(path: Path) -> str | None:
    name = path.name
    for dataset in KNOWN_DATASETS:
        suffix = f"_{dataset}_merged.jsonl"
        if name.endswith(suffix):
            return dataset
    return None


def shorten_sheet_name(name: str) -> str:
    cleaned = name.replace("/", "_")
    for prefix in ("PhilipC_", "ddvd233_", "Qwen_", "keentomato_"):
        if cleaned.startswith(prefix):
            cleaned = cleaned[len(prefix):]
            break
    return cleaned[:31]


def clip_excel_text(value: object) -> str:
    text = "" if value is None else str(value)
    if len(text) <= EXCEL_MAX_CELL_LENGTH:
        return text
    return text[: EXCEL_MAX_CELL_LENGTH - 14] + "\n[TRUNCATED]"


def discover_merged_files(results_dir: Path) -> dict[str, dict[str, Path]]:
    files_by_model: dict[str, dict[str, Path]] = defaultdict(dict)
    for path in sorted(results_dir.glob("*_merged.jsonl")):
        dataset = detect_dataset_from_filename(path)
        if dataset is None:
            continue
        model_slug = path.name[: -len(f"_{dataset}_merged.jsonl")]
        files_by_model[model_slug][dataset] = path
    return dict(files_by_model)


def detect_response_key(row: dict, model_slug: str) -> str:
    expected = f"model_response_{model_slug}"
    if expected in row:
        return expected

    response_keys = [key for key in row if key.startswith("model_response_")]
    if len(response_keys) == 1:
        return response_keys[0]
    if response_keys:
        for key in response_keys:
            if key.endswith(model_slug):
                return key
    raise KeyError(f"Could not determine response key for model slug '{model_slug}'")


def load_rows_by_orig_idx(jsonl_path: Path, model_slug: str) -> dict[int, dict]:
    with jsonl_path.open("r", encoding="utf-8") as f:
        rows = [json.loads(line) for line in f if line.strip()]

    if not rows:
        return {}

    response_key = detect_response_key(rows[0], model_slug)
    indexed_rows: dict[int, dict] = {}

    for row in rows:
        if "_orig_idx" not in row:
            continue
        idx = row["_orig_idx"]
        response = row.get(response_key)
        if not response:
            continue
        indexed_rows[idx] = {
            "_orig_idx": idx,
            "dataset": row.get("dataset", ""),
            "problem": row.get("problem", ""),
            "answer": row.get("answer", ""),
            "sample_id": row.get("sample_id") or row.get("subject_id") or idx,
            "response": response,
        }

    return indexed_rows


def collect_pair_rows(
    anchor_files: dict[str, Path],
    baseline_files: dict[str, Path],
    anchor_slug: str,
    baseline_slug: str,
) -> list[dict]:
    paired_rows: list[dict] = []

    for dataset in KNOWN_DATASETS:
        anchor_path = anchor_files.get(dataset)
        baseline_path = baseline_files.get(dataset)
        if anchor_path is None or baseline_path is None:
            continue

        anchor_rows = load_rows_by_orig_idx(anchor_path, anchor_slug)
        baseline_rows = load_rows_by_orig_idx(baseline_path, baseline_slug)
        shared_indices = sorted(set(anchor_rows) & set(baseline_rows))

        for orig_idx in shared_indices:
            anchor_row = anchor_rows[orig_idx]
            baseline_row = baseline_rows[orig_idx]
            paired_rows.append(
                {
                    "id": f"{dataset}:{anchor_row['sample_id']}",
                    "dataset": dataset,
                    "task": anchor_row["problem"] or baseline_row["problem"],
                    "label": anchor_row["answer"] or baseline_row["answer"],
                    "explanation_a": anchor_row["response"],
                    "explanation_b": baseline_row["response"],
                    "_orig_idx": orig_idx,
                }
            )

    return paired_rows


def sample_rows(rows: list[dict], sample_size: int, seed: int) -> list[dict]:
    if len(rows) <= sample_size:
        return sorted(rows, key=lambda row: (row["dataset"], row["_orig_idx"]))
    rng = random.Random(seed)
    sampled = rng.sample(rows, sample_size)
    return  sorted(sampled, key=lambda row: (row["dataset"], row["_orig_idx"]))


def autosize_worksheet(ws) -> None:
    widths = {
        "A": 22,
        "B": 80,
        "C": 18,
        "D": 90,
        "E": 90,
        "F": 12,
        "G": 12,
        "H": 12,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def set_compiled_worksheet_widths(ws) -> None:
    widths = {
        "A": 28,
        "B": 22,
        "C": 80,
        "D": 18,
        "E": 90,
        "F": 90,
        "G": 12,
        "H": 12,
        "I": 12,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def write_sheet(ws, rows: Iterable[dict], include_baseline_model: bool = False) -> None:
    header = ["ID", "Task", "Label", "Explanation A", "Explanation B", "Specificity", "Coherence", "Concision"]
    if include_baseline_model:
        header = ["Baseline Model", *header]
    ws.append(header)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        values = [
            clip_excel_text(row["id"]),
            clip_excel_text(row["task"]),
            clip_excel_text(row["label"]),
            clip_excel_text(row["explanation_a"]),
            clip_excel_text(row["explanation_b"]),
            "",
            "",
            "",
        ]
        if include_baseline_model:
            values = [clip_excel_text(row["baseline_model"]), *values]
        ws.append(values)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    if include_baseline_model:
        set_compiled_worksheet_widths(ws)
    else:
        autosize_worksheet(ws)


def write_instructions_sheet(ws) -> None:
    lines = [
        "Annotation Instructions",
        "",
        "You will be shown pairs of reasoning explanations (A and B) for the same prediction. Your task is to compare them and select which explanation is better along each criterion.",
        "",
        "For each row, choose:",
        "A if Explanation A is better",
        "B if Explanation B is better",
        "Tie if they are similar",
        "Essentially put A, B or Tie under the Specificity, Coherence, Concision columns respectively",
        "",
        "Evaluation Criteria",
        "Specificity: Which explanation is more concrete and avoids vague or generic statements?",
        "Coherence: Which explanation is more logically structured and easier to follow?",
        "Concision: Which explanation conveys its key points more efficiently without unnecessary verbosity?",
        "",
        "Notes",
        "Focus only on the quality of the reasoning explanation.",
        "Do not consider writing style alone.",
        "The underlying input data is not shown; evaluate based on the explanations themselves.",
        "Work quickly and rely on your first judgment",
    ]

    for line in lines:
        ws.append([line])

    ws["A1"].font = Font(bold=True, size=14)
    for row_idx in range(1, ws.max_row + 1):
        ws[f"A{row_idx}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140


def build_workbook(
    output_path: Path,
    sampled_rows_by_baseline: dict[str, list[dict]],
    metadata_rows: list[dict],
) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "summary"
    summary.append(["Baseline Model", "Available Matched Rows", "Sampled Rows", "Datasets"])

    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in metadata_rows:
        summary.append(
            [
                row["baseline_model"],
                row["available_rows"],
                row["sampled_rows"],
                ", ".join(row["datasets"]),
            ]
        )

    autosize_worksheet(summary)
    summary.freeze_panes = "A2"

    instructions = wb.create_sheet(title="instructions")
    write_instructions_sheet(instructions)

    compiled_rows = []
    for baseline_model, rows in sampled_rows_by_baseline.items():
        for row in rows:
            compiled_rows.append({**row, "baseline_model": baseline_model})

    compiled_rows.sort(key=lambda row: (row["baseline_model"], row["dataset"], row["_orig_idx"]))
    overall = wb.create_sheet(title="overall_compiled")
    write_sheet(overall, compiled_rows, include_baseline_model=True)

    for baseline_model, rows in sampled_rows_by_baseline.items():
        ws = wb.create_sheet(title=shorten_sheet_name(slugify_model(baseline_model)))
        write_sheet(ws, rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def upload_workbook_to_wandb(args: argparse.Namespace, workbook_path: Path) -> None:
    if not getattr(args, "wandb_project", None):
        return

    try:
        import wandb

        run_id = getattr(args, "wandb_run_id", None) or None
        run_name = getattr(args, "wandb_run_name", None) or run_id
        entity = getattr(args, "wandb_entity", None) or None
        artifact_name = (
            getattr(args, "wandb_artifact_name", None)
            or workbook_path.stem.replace(" ", "_")
        )
        artifact_type = getattr(args, "wandb_artifact_type", None) or "annotation_workbook"

        run = wandb.init(
            project=args.wandb_project,
            entity=entity,
            id=run_id,
            name=run_name,
            resume="allow",
            reinit=True,
        )
        artifact = wandb.Artifact(
            name=artifact_name,
            type=artifact_type,
            metadata={
                "anchor_model": ANCHOR_MODEL,
                "baseline_models": BASELINE_MODELS,
                "samples_per_baseline": args.samples_per_baseline,
                "seed": args.seed,
                "file_name": workbook_path.name,
            },
        )
        artifact.add_file(str(workbook_path), name=workbook_path.name)
        run.log_artifact(artifact)
        run.finish()
        print(
            f"W&B: uploaded workbook artifact '{artifact_name}' "
            f"to project '{args.wandb_project}'"
        )
    except Exception as exc:
        print(f"[WARN] W&B artifact upload failed: {exc}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build an Excel workbook of Harpo-vs-baseline reasoning traces."
    )
    parser.add_argument(
        "--results_dir",
        type=Path,
        default=RESULTS_DIR,
        help=f"Directory containing *_merged.jsonl files (default: {RESULTS_DIR})",
    )
    parser.add_argument(
        "--samples_per_baseline",
        type=int,
        default=DEFAULT_SAMPLES_PER_BASELINE,
        help=f"Random matched samples per baseline model (default: {DEFAULT_SAMPLES_PER_BASELINE})",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=DEFAULT_SEED,
        help=f"Random seed for sampling (default: {DEFAULT_SEED})",
    )
    parser.add_argument(
        "--output_xlsx",
        type=Path,
        default=None,
        help="Output workbook path. Default: <results_dir>/reasoning_annotation_workbook.xlsx",
    )
    parser.add_argument(
        "--wandb_project",
        default=WANDB_PROJECT,
        help=f"W&B project name (default: {WANDB_PROJECT!r}; set to '' to disable)",
    )
    parser.add_argument(
        "--wandb_run_id",
        default=WANDB_RUN_ID,
        help=f"W&B run ID (default: {WANDB_RUN_ID!r})",
    )
    parser.add_argument(
        "--wandb_run_name",
        default=WANDB_RUN_NAME,
        help=f"W&B run display name (default: {WANDB_RUN_NAME!r})",
    )
    parser.add_argument(
        "--wandb_entity",
        default=WANDB_ENTITY,
        help="W&B entity (org/team)",
    )
    parser.add_argument(
        "--wandb_artifact_name",
        default=WANDB_ARTIFACT_NAME,
        help=f"Artifact name (default: {WANDB_ARTIFACT_NAME!r})",
    )
    parser.add_argument(
        "--wandb_artifact_type",
        default=WANDB_ARTIFACT_TYPE,
        help=f"W&B artifact type (default: {WANDB_ARTIFACT_TYPE!r})",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    results_dir = args.results_dir.resolve()
    output_xlsx = (
        args.output_xlsx.resolve()
        if args.output_xlsx is not None
        else results_dir / "reasoning_annotation_workbook.xlsx"
    )

    files_by_model = discover_merged_files(results_dir)
    anchor_slug = slugify_model(ANCHOR_MODEL)
    anchor_files = files_by_model.get(anchor_slug)

    if not files_by_model:
        raise FileNotFoundError(f"No merged JSONLs found in {results_dir}")
    if not anchor_files:
        raise FileNotFoundError(
            f"Missing anchor model files for {ANCHOR_MODEL} ({anchor_slug}) in {results_dir}"
        )

    sampled_rows_by_baseline: dict[str, list[dict]] = {}
    metadata_rows: list[dict] = []

    for baseline_model in BASELINE_MODELS:
        baseline_slug = slugify_model(baseline_model)
        baseline_files = files_by_model.get(baseline_slug)
        if not baseline_files:
            print(f"[SKIP] No merged JSONLs found for baseline model: {baseline_model}")
            continue

        paired_rows = collect_pair_rows(anchor_files, baseline_files, anchor_slug, baseline_slug)
        if not paired_rows:
            print(f"[SKIP] No matched rows found for baseline model: {baseline_model}")
            continue

        sampled_rows = sample_rows(paired_rows, args.samples_per_baseline, args.seed)
        sampled_rows_by_baseline[baseline_model] = sampled_rows
        metadata_rows.append(
            {
                "baseline_model": baseline_model,
                "available_rows": len(paired_rows),
                "sampled_rows": len(sampled_rows),
                "datasets": sorted({row["dataset"] for row in paired_rows}),
            }
        )

        print(
            f"[OK] {baseline_model}: sampled {len(sampled_rows)} / {len(paired_rows)} matched rows"
        )

    if not sampled_rows_by_baseline:
        raise RuntimeError("No workbook sheets could be created; check results_dir and model files.")

    build_workbook(output_xlsx, sampled_rows_by_baseline, metadata_rows)
    print(f"Workbook written to: {output_xlsx}")
    upload_workbook_to_wandb(args, output_xlsx)


if __name__ == "__main__":
    main()
